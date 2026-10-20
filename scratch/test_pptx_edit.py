import pptx
from pptx.util import Inches, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation('아이숨(ISOOM) 중간프로젝트 발표_backup.pptx')
print(f"Original slide count: {len(prs.slides)}")

# Test adding a slide
blank_layout = prs.slide_layouts[6]
test_slide = prs.slides.add_slide(blank_layout)
print(f"New slide added, total: {len(prs.slides)}")
